"""Extract every slide's main picture (cropped as on the slide) plus the
slide's own text labels and pointer lines, mapped onto that picture.
Writes pics/<deck>/s<slide>.png and pics/<deck>/index.json."""
import io, json, os, sys
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from scan import walk, is_line, line_ends

Image.MAX_IMAGE_PIXELS = None
deck = sys.argv[1]
prs = Presentation(f'decks/{deck}.pptx')
W, H = prs.slide_width, prs.slide_height
outdir = f'pics/{deck}'
os.makedirs(outdir, exist_ok=True)
index = {}

for n, slide in enumerate(prs.slides, 1):
    pics, texts, lines = [], [], []
    for s, box in walk(slide.shapes):
        img = None
        try:
            if s.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.PLACEHOLDER, MSO_SHAPE_TYPE.LINKED_PICTURE):
                img = s.image
        except Exception:
            img = None
        if img is not None:
            pics.append((s, box, img))
        elif is_line(s):
            lines.append(line_ends(s, box))
        elif s.has_text_frame and s.text_frame.text.strip():
            texts.append((' '.join(s.text_frame.text.split()), box))
    pics = [p for p in pics if p[2].ext.lower() in ('png', 'jpg', 'jpeg', 'gif', 'bmp', 'tif', 'tiff')]
    if not pics:
        continue
    s, (bx, by, bw, bh), img = max(pics, key=lambda p: p[1][2] * p[1][3])
    area = bw * bh / (W * H)
    if area < 0.12:
        continue
    try:
        im = Image.open(io.BytesIO(img.blob))
        im.load()
    except Exception:
        continue
    if im.mode not in ('RGB', 'RGBA', 'L'):
        im = im.convert('RGBA')
    iw, ih = im.size
    cl, cr, ct, cb = s.crop_left or 0, s.crop_right or 0, s.crop_top or 0, s.crop_bottom or 0
    box = (max(0, int(iw * cl)), max(0, int(ih * ct)), min(iw, int(iw * (1 - cr))), min(ih, int(ih * (1 - cb))))
    if box[2] - box[0] < 50 or box[3] - box[1] < 50:
        continue
    im = im.crop(box)
    if im.width < 250 or im.height < 200:
        continue
    # flatten transparency on white
    if im.mode in ('RGBA', 'LA', 'P'):
        bg = Image.new('RGB', im.size, 'white')
        bg.paste(im.convert('RGBA'), mask=im.convert('RGBA').split()[-1])
        im = bg
    else:
        im = im.convert('RGB')
    im.save(f'{outdir}/s{n}.png')

    def rel(x, y):
        return ((x - bx) / bw, (y - by) / bh)

    labels = []
    for t, (tx, ty, tw, th) in texts:
        if len(t) > 50:
            continue
        labels.append({'text': t, 'box': [*rel(tx, ty), tw / bw, th / bh]})
    segs = []
    for (p1, p2, head, tail) in lines:
        segs.append({'a': rel(*p1), 'b': rel(*p2), 'head': head, 'tail': tail})
    index[n] = {'w': im.width, 'h': im.height, 'area': round(area, 2), 'npics': len(pics), 'labels': labels, 'lines': segs,
                'texts': [t for t, _ in texts][:8]}

json.dump(index, open(f'{outdir}/index.json', 'w'), indent=0)
print(deck, len(index))
