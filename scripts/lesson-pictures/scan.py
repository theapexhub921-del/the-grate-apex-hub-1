"""Scan a deck for 'labelled diagram' slides: one main picture plus short
text labels joined to it by lines. Prints JSON per candidate slide."""
import json, sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from lxml import etree

NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}


def walk(shapes, off=(0, 0, 1, 1)):
    """Yield (shape, absolute box) — flattening groups."""
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            # group child coordinates are in the group's child space
            x = s._element.find('.//p:grpSpPr/a:xfrm', NS)
            if x is None:
                yield from walk(s.shapes, off)
                continue
            o = x.find('a:off', NS); e = x.find('a:ext', NS)
            co = x.find('a:chOff', NS); ce = x.find('a:chExt', NS)
            gx, gy, gw, gh = int(o.get('x')), int(o.get('y')), int(e.get('cx')), int(e.get('cy'))
            cx, cy, cw, ch = int(co.get('x')), int(co.get('y')), int(ce.get('cx')) or 1, int(ce.get('cy')) or 1
            sx, sy = gw / cw, gh / ch
            ox, oy, osx, osy = off

            def mk(ox=ox, oy=oy, osx=osx, osy=osy, gx=gx, gy=gy, cx=cx, cy=cy, sx=sx, sy=sy):
                return (ox + (gx - cx * sx) * osx, oy + (gy - cy * sy) * osy, osx * sx, osy * sy)
            yield from walk(s.shapes, mk())
        else:
            if s.left is None:
                continue
            ox, oy, sx, sy = off
            yield s, (ox + s.left * sx, oy + s.top * sy, s.width * sx, s.height * sy)


def line_ends(s, box):
    x, y, w, h = box
    xfrm = s._element.find('.//a:xfrm', NS)
    fh = xfrm is not None and xfrm.get('flipH') == '1'
    fv = xfrm is not None and xfrm.get('flipV') == '1'
    x1, x2 = (x + w, x) if fh else (x, x + w)
    y1, y2 = (y + h, y) if fv else (y, y + h)
    # which end has the arrowhead?
    ln = s._element.find('.//a:ln', NS)
    head = tail = False
    if ln is not None:
        head = ln.find('a:headEnd', NS) is not None and ln.find('a:headEnd', NS).get('type') not in (None, 'none')
        tail = ln.find('a:tailEnd', NS) is not None and ln.find('a:tailEnd', NS).get('type') not in (None, 'none')
    return (x1, y1), (x2, y2), head, tail


def is_line(s):
    el = s._element
    tag = etree.QName(el).localname
    if tag == 'cxnSp':
        return True
    g = el.find('.//a:prstGeom', NS)
    return g is not None and g.get('prst') in ('line', 'straightConnector1', 'bentConnector2', 'bentConnector3', 'curvedConnector3')


def scan(path):
    prs = Presentation(path)
    W, H = prs.slide_width, prs.slide_height
    out = []
    for i, slide in enumerate(prs.slides, 1):
        pics, texts, lines = [], [], []
        for s, box in walk(slide.shapes):
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE or (getattr(s, 'image', None) is not None and s.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and hasattr(s, 'image')):
                pics.append((s, box))
            elif is_line(s):
                lines.append(line_ends(s, box))
            elif s.has_text_frame and s.text_frame.text.strip():
                t = ' '.join(s.text_frame.text.split())
                texts.append((t, box))
        if not pics:
            continue
        big = max(pics, key=lambda p: p[1][2] * p[1][3])
        bx, by, bw, bh = big[1]
        area = bw * bh / (W * H)
        labels = [t for t in texts if len(t[0]) <= 45]
        out.append({'slide': i, 'pics': len(pics), 'area': round(area, 2), 'labels': len(labels), 'lines': len(lines),
                    'title': (texts[0][0][:60] if texts else ''), 'labeltexts': [t[0] for t in labels][:14]})
    return out


if __name__ == '__main__':
    for r in scan(sys.argv[1]):
        if r['labels'] >= 3 and r['area'] > 0.15:
            print(json.dumps(r, ensure_ascii=False))
